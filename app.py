from __future__ import annotations
import csv, io, json, os, secrets, sqlite3, re, threading, webbrowser
from datetime import datetime, timedelta, timezone
from functools import wraps
from pathlib import Path
from flask import Flask, jsonify, render_template, request, session, send_file
from flask_sqlalchemy import SQLAlchemy
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, cast, String
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.middleware.proxy_fix import ProxyFix

BASE_DIR=Path(__file__).resolve().parent; DATA_DIR=BASE_DIR/'data'; DATA_DIR.mkdir(exist_ok=True)
def database_url():
    url=os.environ.get('DATABASE_URL','').strip()
    if not url:
        return 'sqlite:///'+str((DATA_DIR/'vouchers_online.db').resolve()).replace('\\','/')
    # Production uses psycopg v3. Accept common provider URLs and force the
    # SQLAlchemy psycopg dialect so local Windows testing never needs pg_config.
    if url.startswith('postgres://'):
        url='postgresql://'+url[len('postgres://'):]
    if url.startswith('postgresql://'):
        url='postgresql+psycopg://'+url[len('postgresql://'):]
    return url
def secret_key():
    env=os.environ.get('SECRET_KEY','').strip()
    if env: return env
    f=DATA_DIR/'secret.key'
    if f.exists(): return f.read_text(encoding='utf-8').strip()
    value=secrets.token_hex(32); f.write_text(value,encoding='utf-8'); return value
app=Flask(__name__,static_folder='static',template_folder='templates')
app.config.update(SECRET_KEY=secret_key(),SQLALCHEMY_DATABASE_URI=database_url(),SQLALCHEMY_TRACK_MODIFICATIONS=False,MAX_CONTENT_LENGTH=15*1024*1024,SESSION_COOKIE_HTTPONLY=True,SESSION_COOKIE_SAMESITE='Lax',SESSION_COOKIE_SECURE=os.environ.get('COOKIE_SECURE','0')=='1',PERMANENT_SESSION_LIFETIME=timedelta(hours=10))
app.wsgi_app=ProxyFix(app.wsgi_app,x_for=1,x_proto=1,x_host=1)
db=SQLAlchemy(app)
ROLE_LABELS={'admin':'مدير النظام','editor':'إدخال وتعديل','viewer':'مشاهدة فقط'}; WRITE_ROLES={'admin','editor'}

class User(db.Model):
    __tablename__='users'; id=db.Column(db.Integer,primary_key=True); username=db.Column(db.String(80),unique=True,nullable=False,index=True); display_name=db.Column(db.String(150),nullable=False); password_hash=db.Column(db.String(255),nullable=False); role=db.Column(db.String(20),nullable=False,default='viewer'); active=db.Column(db.Boolean,nullable=False,default=True); created_at=db.Column(db.DateTime(timezone=True),nullable=False,default=lambda:datetime.now(timezone.utc)); last_login_at=db.Column(db.DateTime(timezone=True)); session_version=db.Column(db.Integer,nullable=False,default=0)
class Company(db.Model):
    __tablename__='companies'; id=db.Column(db.Integer,primary_key=True); name=db.Column(db.String(220),unique=True,nullable=False,index=True); created_at=db.Column(db.DateTime(timezone=True),nullable=False,default=lambda:datetime.now(timezone.utc))
class Voucher(db.Model):
    __tablename__='vouchers'; id=db.Column(db.Integer,primary_key=True); company_id=db.Column(db.Integer,db.ForeignKey('companies.id',ondelete='RESTRICT'),nullable=False,index=True); receipt_no=db.Column(db.Integer); receipt_date=db.Column(db.Date); amount=db.Column(db.Numeric(18,2)); details=db.Column(db.Text,default=''); created_at=db.Column(db.DateTime(timezone=True),nullable=False,default=lambda:datetime.now(timezone.utc)); updated_at=db.Column(db.DateTime(timezone=True),nullable=False,default=lambda:datetime.now(timezone.utc),onupdate=lambda:datetime.now(timezone.utc)); created_by_id=db.Column(db.Integer,db.ForeignKey('users.id')); updated_by_id=db.Column(db.Integer,db.ForeignKey('users.id')); company=db.relationship('Company'); created_by=db.relationship('User',foreign_keys=[created_by_id]); updated_by=db.relationship('User',foreign_keys=[updated_by_id]); __table_args__=(db.UniqueConstraint('company_id','receipt_no',name='uq_company_receipt'),)
class AuditLog(db.Model):
    __tablename__='audit_log'; id=db.Column(db.Integer,primary_key=True); user_id=db.Column(db.Integer,db.ForeignKey('users.id')); action=db.Column(db.String(80),nullable=False); entity_type=db.Column(db.String(40),nullable=False); entity_id=db.Column(db.Integer); details=db.Column(db.Text,default=''); created_at=db.Column(db.DateTime(timezone=True),nullable=False,default=lambda:datetime.now(timezone.utc),index=True); user=db.relationship('User')

def iso_date(v): return v.isoformat() if v else None
def iso_dt(v): return v.isoformat(timespec='seconds') if v else None
def user_dict(u): return {'id':u.id,'username':u.username,'display_name':u.display_name,'role':u.role,'role_label':ROLE_LABELS.get(u.role,u.role),'active':bool(u.active),'created_at':iso_dt(u.created_at),'last_login_at':iso_dt(u.last_login_at)}
def current_user():
    uid=session.get('user_id'); u=db.session.get(User,uid) if uid else None; return u if u and u.active and session.get('session_version')==u.session_version else None
def ensure_csrf():
    if 'csrf_token' not in session: session['csrf_token']=secrets.token_urlsafe(24)
    return session['csrf_token']
def login_required(fn):
    @wraps(fn)
    def w(*a,**k):
        if not current_user(): return jsonify(ok=False,message='يرجى تسجيل الدخول.'),401
        return fn(*a,**k)
    return w
def roles_required(*roles):
    def deco(fn):
        @wraps(fn)
        @login_required
        def w(*a,**k):
            if current_user().role not in roles: return jsonify(ok=False,message='ليس لديك صلاحية لتنفيذ هذه العملية.'),403
            return fn(*a,**k)
        return w
    return deco
@app.before_request
def csrf_guard():
    if request.method in {'POST','PUT','PATCH','DELETE'} and request.path.startswith('/api/') and request.path!='/api/login' and current_user():
        token=request.headers.get('X-CSRF-Token','')
        if not token or not secrets.compare_digest(token,session.get('csrf_token','')): return jsonify(ok=False,message='انتهت صلاحية نموذج الحماية. أعد تحميل الصفحة.'),403
def audit(action,entity_type,entity_id=None,details=''):
    u=current_user(); db.session.add(AuditLog(user_id=u.id if u else None,action=action,entity_type=entity_type,entity_id=entity_id,details=details[:1000]))
def parse_int_optional(v):
    if v is None or str(v).strip()=='': return None
    s=str(v).strip()
    if not s.isdigit(): raise ValueError('رقم الوصل يجب أن يكون رقماً صحيحاً.')
    return int(s)
def parse_amount_optional(v):
    if v is None or str(v).strip()=='': return None
    try: return round(float(str(v).replace(',','').strip()),2)
    except: raise ValueError('المبلغ يجب أن يكون رقماً.')
def parse_date_optional(v):
    if v is None or str(v).strip()=='': return None
    try: return datetime.strptime(str(v).strip(),'%Y-%m-%d').date()
    except: raise ValueError('صيغة التاريخ غير صحيحة.')
def validate_voucher(company_id,receipt_no,receipt_date,exclude_id=None):
    if not db.session.get(Company,company_id): raise ValueError('الشركة المحددة غير موجودة.')
    if receipt_no is None:
        if receipt_date: raise ValueError('لا يمكن إدخال تاريخ بدون رقم وصل.')
        return
    q=Voucher.query.filter_by(company_id=company_id,receipt_no=receipt_no)
    if exclude_id: q=q.filter(Voucher.id!=exclude_id)
    if q.first(): raise ValueError('رقم الوصل مكرر لنفس الشركة.')
    if not receipt_date: return
    prev=Voucher.query.filter(Voucher.company_id==company_id,Voucher.receipt_no<receipt_no,Voucher.receipt_date>=receipt_date); nxt=Voucher.query.filter(Voucher.company_id==company_id,Voucher.receipt_no>receipt_no,Voucher.receipt_date<=receipt_date)
    if exclude_id: prev=prev.filter(Voucher.id!=exclude_id); nxt=nxt.filter(Voucher.id!=exclude_id)
    bp=prev.order_by(Voucher.receipt_no.desc()).first(); bn=nxt.order_by(Voucher.receipt_no.asc()).first()
    if bp: raise ValueError(f'التاريخ غير متوافق: الوصل الأصغر رقم {bp.receipt_no} تاريخه {iso_date(bp.receipt_date)} ويجب أن يكون أقدم.')
    if bn: raise ValueError(f'التاريخ غير متوافق: الوصل الأكبر رقم {bn.receipt_no} تاريخه {iso_date(bn.receipt_date)} ويجب أن يكون أحدث.')
def voucher_dict(v): return {'id':v.id,'company_id':v.company_id,'company_name':v.company.name,'receipt_no':v.receipt_no,'receipt_date':iso_date(v.receipt_date),'amount':float(v.amount) if v.amount is not None else None,'details':v.details or '','created_at':iso_dt(v.created_at),'updated_at':iso_dt(v.updated_at),'created_by':v.created_by.display_name if v.created_by else None,'updated_by':v.updated_by.display_name if v.updated_by else None}
def voucher_query_from_args(args):
    q=Voucher.query.join(Company); cid=args.get('company_id','').strip()
    if cid.isdigit(): q=q.filter(Voucher.company_id==int(cid))
    if args.get('date_from'): q=q.filter(Voucher.receipt_date>=parse_date_optional(args.get('date_from')))
    if args.get('date_to'): q=q.filter(Voucher.receipt_date<=parse_date_optional(args.get('date_to')))
    if args.get('receipt_from','').isdigit(): q=q.filter(Voucher.receipt_no>=int(args.get('receipt_from')))
    if args.get('receipt_to','').isdigit(): q=q.filter(Voucher.receipt_no<=int(args.get('receipt_to')))
    s=args.get('q','').strip()
    if s:
        like=f'%{s}%'; q=q.filter(or_(Company.name.ilike(like),Voucher.details.ilike(like),cast(Voucher.receipt_no,String).ilike(like)))
    return q

@app.errorhandler(ValueError)
def value_error(e): return jsonify(ok=False,message=str(e)),400
@app.errorhandler(413)
def too_large(e): return jsonify(ok=False,message='حجم الملف أكبر من الحد المسموح.'),413
@app.route('/')
def index(): return render_template('index.html')
@app.get('/health')
def health(): return jsonify(ok=True,status='healthy')

@app.post('/api/login')
def api_login():
    d=request.get_json(silent=True) or {}; username=str(d.get('username','')).strip(); password=str(d.get('password',''))
    u=User.query.filter(func.lower(User.username)==username.lower()).first()
    if not u or not u.active or not check_password_hash(u.password_hash,password): return jsonify(ok=False,message='اسم المستخدم أو كلمة المرور غير صحيحة.'),401
    session.clear(); session.permanent=True; session['user_id']=u.id; session['session_version']=u.session_version; ensure_csrf(); u.last_login_at=datetime.now(timezone.utc); db.session.commit()
    return jsonify(ok=True,message='تم تسجيل الدخول بنجاح.',user=user_dict(u),csrf_token=session['csrf_token'])
@app.post('/api/logout')
@login_required
def api_logout(): session.clear(); return jsonify(ok=True)
@app.get('/api/session')
def api_session():
    u=current_user(); return jsonify(authenticated=False) if not u else jsonify(authenticated=True,user=user_dict(u),csrf_token=ensure_csrf())
@app.post('/api/change-password')
@login_required
def api_change_password():
    u=current_user(); d=request.get_json(silent=True) or {}; cur=str(d.get('current_password','')); new=str(d.get('new_password',''))
    if not check_password_hash(u.password_hash,cur): return jsonify(ok=False,message='كلمة المرور الحالية غير صحيحة.'),400
    if len(new)<8: return jsonify(ok=False,message='كلمة المرور الجديدة يجب أن تكون 8 أحرف على الأقل.'),400
    if new==cur: return jsonify(ok=False,message='اختر كلمة مرور جديدة مختلفة.'),400
    u.password_hash=generate_password_hash(new,method='pbkdf2:sha256:600000'); u.session_version+=1; audit('change_password','user',u.id,'تغيير كلمة المرور من المستخدم'); db.session.commit(); session.clear()
    return jsonify(ok=True,message='تم تغيير كلمة المرور. سجّل الدخول مرة أخرى.')

@app.get('/api/users')
@roles_required('admin')
def api_users(): return jsonify([user_dict(u) for u in User.query.order_by(User.id).all()])
@app.post('/api/users')
@roles_required('admin')
def api_create_user():
    d=request.get_json(silent=True) or {}; username=str(d.get('username','')).strip(); display=str(d.get('display_name','')).strip(); role=str(d.get('role','viewer')); password=str(d.get('password',''))
    if not re.fullmatch(r'[A-Za-z0-9_.-]{3,40}',username): raise ValueError('اسم المستخدم يجب أن يكون 3-40 حرفاً من الإنجليزية أو الأرقام أو . _ - فقط.')
    if not display: raise ValueError('يرجى إدخال الاسم الظاهر.')
    if role not in ROLE_LABELS: raise ValueError('نوع الصلاحية غير صحيح.')
    if len(password)<8: raise ValueError('كلمة المرور يجب أن تكون 8 أحرف على الأقل.')
    if User.query.filter(func.lower(User.username)==username.lower()).first(): raise ValueError('اسم المستخدم موجود مسبقاً.')
    u=User(username=username,display_name=display,role=role,password_hash=generate_password_hash(password,method='pbkdf2:sha256:600000')); db.session.add(u); db.session.flush(); audit('create','user',u.id,f'{username} / {role}'); db.session.commit()
    return jsonify(ok=True,message='تم إنشاء المستخدم.',user=user_dict(u)),201
@app.put('/api/users/<int:uid>')
@roles_required('admin')
def api_update_user(uid):
    u=db.session.get(User,uid)
    if not u: return jsonify(ok=False,message='المستخدم غير موجود.'),404
    d=request.get_json(silent=True) or {}; display=str(d.get('display_name',u.display_name)).strip(); role=str(d.get('role',u.role)); active=bool(d.get('active',u.active))
    if not display or role not in ROLE_LABELS: raise ValueError('بيانات المستخدم غير صحيحة.')
    me=current_user()
    if u.id==me.id and not active: raise ValueError('لا يمكنك تعطيل حسابك الحالي.')
    if u.role=='admin' and (role!='admin' or not active):
        if User.query.filter_by(role='admin',active=True).count()<=1: raise ValueError('يجب أن يبقى مدير نظام فعّال واحد على الأقل.')
    u.display_name=display; u.role=role; u.active=active; audit('update','user',u.id,f'role={role}, active={active}'); db.session.commit(); return jsonify(ok=True,message='تم تحديث المستخدم.',user=user_dict(u))
@app.post('/api/users/<int:uid>/reset-password')
@roles_required('admin')
def api_reset_password(uid):
    u=db.session.get(User,uid)
    if not u: return jsonify(ok=False,message='المستخدم غير موجود.'),404
    new=str((request.get_json(silent=True) or {}).get('new_password',''))
    if len(new)<8: raise ValueError('كلمة المرور الجديدة يجب أن تكون 8 أحرف على الأقل.')
    u.password_hash=generate_password_hash(new,method='pbkdf2:sha256:600000'); u.session_version+=1; audit('reset_password','user',u.id,f'إعادة تعيين بواسطة {current_user().username}'); db.session.commit(); return jsonify(ok=True,message='تم تعيين كلمة المرور الجديدة.')

@app.get('/api/dashboard')
@login_required
def api_dashboard():
    stats={'companies_count':Company.query.count(),'entries_count':Voucher.query.count(),'vouchers_count':Voucher.query.filter(Voucher.receipt_no.isnot(None)).count(),'total_amount':float(db.session.query(func.coalesce(func.sum(Voucher.amount),0)).scalar() or 0)}
    recent=[voucher_dict(v) for v in Voucher.query.order_by(Voucher.id.desc()).limit(8).all()]
    top=db.session.query(Company.id,Company.name,func.coalesce(func.sum(Voucher.amount),0).label('total')).outerjoin(Voucher).group_by(Company.id).order_by(func.coalesce(func.sum(Voucher.amount),0).desc()).limit(10).all()
    return jsonify(stats=stats,recent=recent,top=[{'id':x.id,'name':x.name,'total':float(x.total or 0)} for x in top])
@app.get('/api/companies')
@login_required
def api_companies():
    rs=db.session.query(Company.id,Company.name,Company.created_at,func.count(Voucher.id).label('entries_count'),func.coalesce(func.sum(Voucher.amount),0).label('total_amount')).outerjoin(Voucher).group_by(Company.id).order_by(Company.name).all()
    return jsonify([{'id':r.id,'name':r.name,'created_at':iso_dt(r.created_at),'entries_count':r.entries_count,'total_amount':float(r.total_amount or 0)} for r in rs])
@app.post('/api/companies')
@roles_required('admin','editor')
def api_company_create():
    name=str((request.get_json(silent=True) or {}).get('name','')).strip()
    if not name: raise ValueError('يرجى إدخال اسم الشركة.')
    if Company.query.filter(func.lower(Company.name)==name.lower()).first(): raise ValueError('اسم الشركة موجود مسبقاً.')
    c=Company(name=name); db.session.add(c); db.session.flush(); audit('create','company',c.id,name); db.session.commit(); return jsonify(ok=True,message='تمت إضافة الشركة.',id=c.id),201
@app.put('/api/companies/<int:cid>')
@roles_required('admin','editor')
def api_company_update(cid):
    c=db.session.get(Company,cid)
    if not c: return jsonify(ok=False,message='الشركة غير موجودة.'),404
    name=str((request.get_json(silent=True) or {}).get('name','')).strip()
    if not name: raise ValueError('اسم الشركة لا يمكن أن يكون فارغاً.')
    other=Company.query.filter(func.lower(Company.name)==name.lower(),Company.id!=cid).first()
    if other: raise ValueError('يوجد اسم شركة مطابق مسبقاً.')
    c.name=name; audit('update','company',cid,name); db.session.commit(); return jsonify(ok=True,message='تم تعديل اسم الشركة.')
@app.delete('/api/companies/<int:cid>')
@roles_required('admin','editor')
def api_company_delete(cid):
    c=db.session.get(Company,cid)
    if not c: return jsonify(ok=False,message='الشركة غير موجودة.'),404
    if Voucher.query.filter_by(company_id=cid).count(): raise ValueError('لا يمكن حذف الشركة لأنها مرتبطة بوصولات.')
    audit('delete','company',cid,c.name); db.session.delete(c); db.session.commit(); return jsonify(ok=True,message='تم حذف الشركة.')

@app.get('/api/vouchers')
@login_required
def api_vouchers():
    q=voucher_query_from_args(request.args); count=q.count(); total=float(q.with_entities(func.coalesce(func.sum(Voucher.amount),0)).scalar() or 0); rows=q.order_by(Voucher.receipt_date.desc().nullslast(),Voucher.id.desc()).limit(3000).all(); return jsonify(rows=[voucher_dict(v) for v in rows],totals={'count_rows':count,'total_amount':total})
@app.get('/api/vouchers/<int:vid>')
@login_required
def api_voucher_one(vid):
    v=db.session.get(Voucher,vid); return (jsonify(voucher_dict(v)) if v else (jsonify(ok=False,message='السجل غير موجود.'),404))
@app.get('/api/check')
@login_required
def api_check():
    validate_voucher(int(request.args.get('company_id')),parse_int_optional(request.args.get('receipt_no')),parse_date_optional(request.args.get('receipt_date')),int(request.args.get('exclude_id')) if request.args.get('exclude_id','').isdigit() else None); return jsonify(ok=True,message='البيانات متوافقة مع قواعد الوصولات.')
@app.post('/api/vouchers')
@roles_required('admin','editor')
def api_voucher_create():
    d=request.get_json(silent=True) or {}; cid=int(d.get('company_id')); no=parse_int_optional(d.get('receipt_no')); dt=parse_date_optional(d.get('receipt_date')); amount=parse_amount_optional(d.get('amount')); details=str(d.get('details') or '').strip(); validate_voucher(cid,no,dt); u=current_user(); v=Voucher(company_id=cid,receipt_no=no,receipt_date=dt,amount=amount,details=details,created_by_id=u.id,updated_by_id=u.id); db.session.add(v); db.session.flush(); audit('create','voucher',v.id,f'company={cid}, receipt={no}'); db.session.commit(); return jsonify(ok=True,message='تم حفظ الوصل بنجاح.',id=v.id),201
@app.put('/api/vouchers/<int:vid>')
@roles_required('admin','editor')
def api_voucher_update(vid):
    v=db.session.get(Voucher,vid)
    if not v: return jsonify(ok=False,message='السجل غير موجود.'),404
    d=request.get_json(silent=True) or {}; cid=int(d.get('company_id')); no=parse_int_optional(d.get('receipt_no')); dt=parse_date_optional(d.get('receipt_date')); amount=parse_amount_optional(d.get('amount')); details=str(d.get('details') or '').strip(); validate_voucher(cid,no,dt,vid); v.company_id=cid; v.receipt_no=no; v.receipt_date=dt; v.amount=amount; v.details=details; v.updated_by_id=current_user().id; audit('update','voucher',vid,f'company={cid}, receipt={no}'); db.session.commit(); return jsonify(ok=True,message='تم تحديث الوصل بنجاح.')
@app.delete('/api/vouchers/<int:vid>')
@roles_required('admin','editor')
def api_voucher_delete(vid):
    v=db.session.get(Voucher,vid)
    if not v: return jsonify(ok=False,message='السجل غير موجود.'),404
    audit('delete','voucher',vid,f'receipt={v.receipt_no}'); db.session.delete(v); db.session.commit(); return jsonify(ok=True,message='تم حذف الوصل.')

@app.get('/api/reports')
@login_required
def api_reports():
    dfrom=parse_date_optional(request.args.get('date_from')); dto=parse_date_optional(request.args.get('date_to')); rows=[]; grand=0
    for c in Company.query.order_by(Company.name).all():
        q=Voucher.query.filter_by(company_id=c.id)
        if dfrom: q=q.filter(Voucher.receipt_date>=dfrom)
        if dto: q=q.filter(Voucher.receipt_date<=dto)
        vs=q.all(); total=sum(float(v.amount or 0) for v in vs); grand+=total; dates=[v.receipt_date for v in vs if v.receipt_date]
        rows.append({'id':c.id,'name':c.name,'entries_count':len(vs),'vouchers_count':sum(1 for v in vs if v.receipt_no is not None),'total_amount':total,'first_date':iso_date(min(dates)) if dates else None,'last_date':iso_date(max(dates)) if dates else None})
    rows.sort(key=lambda x:(-x['total_amount'],x['name'])); return jsonify(summary=rows,grand_total=grand)
@app.get('/api/monthly-trend')
@login_required
def api_monthly_trend():
    dfrom=parse_date_optional(request.args.get('date_from')); dto=parse_date_optional(request.args.get('date_to')); q=Voucher.query.filter(Voucher.receipt_date.isnot(None))
    if dfrom: q=q.filter(Voucher.receipt_date>=dfrom)
    if dto: q=q.filter(Voucher.receipt_date<=dto)
    agg={}
    for v in q.all():
        key=v.receipt_date.strftime('%Y-%m'); agg[key]=agg.get(key,0)+float(v.amount or 0)
    return jsonify([{'month':k,'total':agg[k]} for k in sorted(agg)])

@app.get('/api/company-report/<int:cid>')
@login_required
def api_company_report(cid):
    c=db.session.get(Company,cid)
    if not c: return jsonify(ok=False,message='الشركة غير موجودة.'),404
    dfrom=parse_date_optional(request.args.get('date_from')); dto=parse_date_optional(request.args.get('date_to')); q=Voucher.query.filter_by(company_id=cid)
    if dfrom:q=q.filter(Voucher.receipt_date>=dfrom)
    if dto:q=q.filter(Voucher.receipt_date<=dto)
    vs=q.order_by(Voucher.receipt_no.asc().nullslast(),Voucher.id).all(); total=sum(float(v.amount or 0) for v in vs)
    return jsonify(company={'id':c.id,'name':c.name},rows=[voucher_dict(v) for v in vs],total=total,numbered_count=sum(1 for v in vs if v.receipt_no is not None),date_from=iso_date(dfrom),date_to=iso_date(dto))

@app.get('/export.csv')
@login_required
def export_csv():
    q=voucher_query_from_args(request.args).order_by(Company.name,Voucher.receipt_no,Voucher.receipt_date); sio=io.StringIO(); w=csv.writer(sio); w.writerow(['اسم الشركة','رقم الوصل','التاريخ','المبلغ','التفاصيل','أنشأ بواسطة','عدّل بواسطة'])
    for v in q.all(): w.writerow([v.company.name,v.receipt_no or '',iso_date(v.receipt_date) or '',float(v.amount) if v.amount is not None else '',v.details or '',v.created_by.display_name if v.created_by else '',v.updated_by.display_name if v.updated_by else ''])
    bio=io.BytesIO(('\ufeff'+sio.getvalue()).encode('utf-8')); return send_file(bio,mimetype='text/csv',as_attachment=True,download_name='karbala_audit_report.csv')
@app.get('/backup.xlsx')
@login_required
def backup_xlsx():
    wb=Workbook(); ws=wb.active; ws.title='الوصولات'; ws.append(['اسم الشركة','رقم الوصل','التاريخ','المبلغ','التفاصيل','أنشأ بواسطة','عدّل بواسطة'])
    for v in Voucher.query.join(Company).order_by(Company.name,Voucher.id).all(): ws.append([v.company.name,v.receipt_no,iso_date(v.receipt_date),float(v.amount) if v.amount is not None else None,v.details or '',v.created_by.display_name if v.created_by else '',v.updated_by.display_name if v.updated_by else ''])
    users=wb.create_sheet('المستخدمون'); users.append(['اسم المستخدم','الاسم الظاهر','الصلاحية','فعال'])
    if current_user().role=='admin':
        for u in User.query.order_by(User.id): users.append([u.username,u.display_name,ROLE_LABELS.get(u.role,u.role),'نعم' if u.active else 'لا'])
    bio=io.BytesIO(); wb.save(bio); bio.seek(0); return send_file(bio,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',as_attachment=True,download_name='karbala_audit_backup.xlsx')
@app.get('/backup.db')
@roles_required('admin')
def backup_db():
    payload={'system':'Karbala Audit Online Backup','version':3,'exported_at':datetime.now(timezone.utc).isoformat(),'companies':[{'id':c.id,'name':c.name,'created_at':iso_dt(c.created_at)} for c in Company.query.order_by(Company.id)],'vouchers':[voucher_dict(v) for v in Voucher.query.order_by(Voucher.id)],'users':[{'username':u.username,'display_name':u.display_name,'role':u.role,'active':u.active} for u in User.query.order_by(User.id)]}
    bio=io.BytesIO(json.dumps(payload,ensure_ascii=False,indent=2).encode('utf-8')); return send_file(bio,mimetype='application/octet-stream',as_attachment=True,download_name='karbala_audit_backup.db')
@app.post('/api/import-xlsx')
@roles_required('admin','editor')
def import_xlsx():
    f=request.files.get('excel_file')
    if not f or not f.filename.lower().endswith('.xlsx'): raise ValueError('اختر ملف Excel بصيغة XLSX.')
    wb=load_workbook(f,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]
    all_rows=list(ws.iter_rows(values_only=True))
    if not all_rows: raise ValueError('ملف Excel فارغ.')
    headers=[str(x or '').strip() for x in all_rows[0]]
    def find_header(*names):
        for i,h in enumerate(headers):
            if any(n in h for n in names): return i
        return None
    ci=find_header('اسم الشركة','الشركة'); ni=find_header('رقم الوصل','رقم السند'); di=find_header('التاريخ'); ai=find_header('المبلغ'); ti=find_header('التفاصيل','البيان')
    if ci is None:
        # توافق مع ملف Excel الأصلي: B الشركة، C الرقم، D التاريخ، E المبلغ، F التفاصيل.
        ci,ni,di,ai,ti=1,2,3,4,5
    imported=skipped=0; errors=[]
    for idx,row in enumerate(all_rows[1:],start=2):
        def get(i): return row[i] if i is not None and i<len(row) else None
        name=str(get(ci) or '').strip()
        if not name: skipped+=1; continue
        c=Company.query.filter(func.lower(Company.name)==name.lower()).first()
        if not c: c=Company(name=name); db.session.add(c); db.session.flush()
        try:
            no=parse_int_optional(get(ni)); rawdt=get(di); dt=rawdt.date() if isinstance(rawdt,datetime) else parse_date_optional(rawdt); amount=parse_amount_optional(get(ai)); details=str(get(ti) or '')
            validate_voucher(c.id,no,dt); u=current_user(); db.session.add(Voucher(company_id=c.id,receipt_no=no,receipt_date=dt,amount=amount,details=details,created_by_id=u.id,updated_by_id=u.id)); imported+=1
        except Exception as e:
            skipped+=1
            if len(errors)<5: errors.append(f'صف {idx}: {e}')
    audit('import','xlsx',None,f'imported={imported}, skipped={skipped}'); db.session.commit(); return jsonify(ok=True,message=f'تم استيراد {imported} سجل وتجاوز {skipped} سجل.',errors=errors)
@app.post('/api/restore-db')
@roles_required('admin')
def restore_db():
    f=request.files.get('backup_file')
    if not f: raise ValueError('اختر ملف DB.')
    try: obj=json.loads(f.read().decode('utf-8-sig'))
    except: raise ValueError('ملف DB غير صحيح.')
    if not isinstance(obj.get('companies'),list) or not isinstance(obj.get('vouchers'),list): raise ValueError('ملف النسخة الاحتياطية غير متوافق.')
    Voucher.query.delete(); Company.query.delete(); db.session.flush(); cmap={}
    for c in obj['companies']:
        nc=Company(name=str(c.get('name','')).strip()); db.session.add(nc); db.session.flush(); cmap[c.get('id')]=nc.id
    u=current_user()
    for v in obj['vouchers']:
        cid=cmap.get(v.get('company_id'))
        if not cid: continue
        db.session.add(Voucher(company_id=cid,receipt_no=parse_int_optional(v.get('receipt_no')),receipt_date=parse_date_optional(v.get('receipt_date')),amount=parse_amount_optional(v.get('amount')),details=str(v.get('details') or ''),created_by_id=u.id,updated_by_id=u.id))
    audit('restore','db',None,f"companies={len(obj['companies'])}, vouchers={len(obj['vouchers'])}"); db.session.commit(); return jsonify(ok=True,message='تمت استعادة بيانات الشركات والوصولات من ملف DB بنجاح.')
@app.get('/api/audit-log')
@roles_required('admin')
def audit_log():
    logs=AuditLog.query.order_by(AuditLog.id.desc()).limit(300).all(); return jsonify([{'id':x.id,'user':x.user.display_name if x.user else 'النظام','action':x.action,'entity_type':x.entity_type,'entity_id':x.entity_id,'details':x.details,'created_at':iso_dt(x.created_at)} for x in logs])

def init_database():
    db.create_all()
    if User.query.count()==0:
        username=os.environ.get('ADMIN_USERNAME','admin').strip() or 'admin'; password=os.environ.get('ADMIN_INITIAL_PASSWORD','12345678'); display=os.environ.get('ADMIN_DISPLAY_NAME','مدير النظام')
        db.session.add(User(username=username,display_name=display,role='admin',active=True,password_hash=generate_password_hash(password,method='pbkdf2:sha256:600000'))); db.session.commit()
    if Company.query.count()==0:
        for name in ['بغداد','الاجيال','المعاصرة','الاهلي']: db.session.add(Company(name=name))
        db.session.commit()
        if Voucher.query.count()==0:
            admin=User.query.filter_by(role='admin').first(); bag=Company.query.filter_by(name='بغداد').first(); aj=Company.query.filter_by(name='الاجيال').first(); db.session.add_all([Voucher(company_id=bag.id,amount=100,created_by_id=admin.id,updated_by_id=admin.id),Voucher(company_id=bag.id,amount=1200,created_by_id=admin.id,updated_by_id=admin.id),Voucher(company_id=aj.id,amount=150,created_by_id=admin.id,updated_by_id=admin.id)]); db.session.commit()
with app.app_context(): init_database()
if __name__=='__main__':
    port=int(os.environ.get('PORT','5000'))
    # In local Windows testing, open the browser automatically. Production
    # servers import app:app and never execute this block.
    if not os.environ.get('DATABASE_URL','').strip():
        threading.Timer(1.2,lambda:webbrowser.open(f'http://127.0.0.1:{port}')).start()
    app.run(host='0.0.0.0',port=port,debug=False)
